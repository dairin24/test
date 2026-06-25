"""要求ファイルからテキスト/ドキュメントを抽出する。

対応形式: Markdown/テキスト, CSV/Excel, Word(.docx), PDF。

PDF は Claude の document 入力 (base64) として直接渡すのが最も精度が高いため、
``ExtractResult`` で「テキスト」か「PDF base64」のどちらかを返す。
"""

from __future__ import annotations

import base64
import io
from dataclasses import dataclass
from typing import Optional

TEXT_EXTS = {".md", ".markdown", ".txt"}
CSV_EXTS = {".csv"}
EXCEL_EXTS = {".xlsx", ".xls"}
WORD_EXTS = {".docx"}
PDF_EXTS = {".pdf"}

SUPPORTED_EXTS = TEXT_EXTS | CSV_EXTS | EXCEL_EXTS | WORD_EXTS | PDF_EXTS


@dataclass
class ExtractResult:
    """抽出結果。

    - ``kind == "text"``: ``text`` に抽出済みテキストが入る。
    - ``kind == "pdf"``: ``pdf_base64`` に PDF の base64 文字列が入る
      (Claude の document ブロックへそのまま渡す)。
    """

    kind: str  # "text" | "pdf"
    text: Optional[str] = None
    pdf_base64: Optional[str] = None


class UnsupportedFileError(ValueError):
    """未対応のファイル形式。"""


def _ext(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot != -1 else ""


def extract(filename: str, data: bytes) -> ExtractResult:
    """ファイル名と中身 (bytes) から ``ExtractResult`` を返す。"""
    ext = _ext(filename)

    if ext in TEXT_EXTS:
        return ExtractResult(kind="text", text=_decode_text(data))
    if ext in CSV_EXTS:
        return ExtractResult(kind="text", text=_extract_csv(data))
    if ext in EXCEL_EXTS:
        return ExtractResult(kind="text", text=_extract_excel(data))
    if ext in WORD_EXTS:
        return ExtractResult(kind="text", text=_extract_docx(data))
    if ext in PDF_EXTS:
        return ExtractResult(
            kind="pdf",
            pdf_base64=base64.standard_b64encode(data).decode("ascii"),
        )

    raise UnsupportedFileError(
        f"未対応のファイル形式です: '{ext or filename}'. "
        f"対応形式: {', '.join(sorted(SUPPORTED_EXTS))}"
    )


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp932", "shift_jis"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    # 最後の手段: 置換しつつデコード
    return data.decode("utf-8", errors="replace")


def _extract_csv(data: bytes) -> str:
    import pandas as pd

    df = pd.read_csv(io.BytesIO(data))
    return _df_to_markdown(df)


def _extract_excel(data: bytes) -> str:
    import pandas as pd

    sheets = pd.read_excel(io.BytesIO(data), sheet_name=None)
    parts: list[str] = []
    for name, df in sheets.items():
        parts.append(f"## シート: {name}\n\n{_df_to_markdown(df)}")
    return "\n\n".join(parts)


def _df_to_markdown(df) -> str:
    """DataFrame を LLM が読みやすい Markdown 表へ変換する。"""
    df = df.fillna("")
    try:
        return df.to_markdown(index=False)
    except Exception:
        # tabulate 未インストール等のフォールバック
        return df.to_csv(index=False)


def _extract_docx(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]

    # 表も抽出する
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(cells))

    return "\n".join(parts)


def extract_pdf_text(data: bytes) -> str:
    """PDF からテキストを抽出する (フォールバック用)。"""
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)
