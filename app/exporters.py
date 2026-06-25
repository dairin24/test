"""要件定義を Markdown / Excel へエクスポートする。"""

from __future__ import annotations

import io

from .schema import RequirementsSpec

# 表の列定義: (属性名, 見出し)
COLUMNS = [
    ("id", "ID"),
    ("title", "タイトル"),
    ("type", "種別"),
    ("asil", "ASIL"),
    ("priority", "優先度"),
    ("description", "内容"),
    ("rationale", "根拠"),
    ("source", "出典"),
    ("verification_method", "検証方法"),
    ("acceptance_criteria", "受入基準"),
]


def _criteria_to_str(value, sep: str) -> str:
    if isinstance(value, list):
        return sep.join(str(v) for v in value)
    return str(value)


def to_markdown(spec: RequirementsSpec) -> str:
    """要件定義を Markdown 文書へ変換する。"""
    lines: list[str] = []
    lines.append("# 要件定義書\n")
    lines.append("## プロジェクト概要\n")
    lines.append(spec.project_overview + "\n")

    lines.append("## 要件一覧\n")
    headers = [h for _, h in COLUMNS]
    lines.append("| " + " | ".join(headers) + " |")
    lines.append("| " + " | ".join("---" for _ in headers) + " |")

    for req in spec.requirements:
        cells = []
        for attr, _ in COLUMNS:
            value = getattr(req, attr)
            if attr == "acceptance_criteria":
                value = _criteria_to_str(value, "<br>")
            # Markdown の表内で改行・パイプを壊さないようにエスケープ
            text = str(value).replace("|", "\\|").replace("\n", "<br>")
            cells.append(text)
        lines.append("| " + " | ".join(cells) + " |")

    lines.append("\n## オープンイシュー / 確認事項\n")
    if spec.open_issues:
        for issue in spec.open_issues:
            lines.append(f"- {issue}")
    else:
        lines.append("- なし")

    return "\n".join(lines) + "\n"


def to_excel(spec: RequirementsSpec) -> bytes:
    """要件定義を .xlsx (bytes) へ変換する。"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()

    # --- 要件一覧シート ---
    ws = wb.active
    ws.title = "要件一覧"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    wrap = Alignment(vertical="top", wrap_text=True)

    headers = [h for _, h in COLUMNS]
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for req in spec.requirements:
        row = []
        for attr, _ in COLUMNS:
            value = getattr(req, attr)
            if attr == "acceptance_criteria":
                value = _criteria_to_str(value, "\n")
            row.append(str(value))
        ws.append(row)

    # 列幅と折り返し
    widths = {
        "ID": 14, "タイトル": 24, "種別": 14, "ASIL": 8, "優先度": 10,
        "内容": 40, "根拠": 30, "出典": 24, "検証方法": 12, "受入基準": 40,
    }
    for col_idx, (_, header) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = widths.get(header, 20)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap
    ws.freeze_panes = "A2"

    # --- 概要シート ---
    ws2 = wb.create_sheet("概要")
    ws2["A1"] = "プロジェクト概要"
    ws2["A1"].font = Font(bold=True)
    ws2["A2"] = spec.project_overview
    ws2["A2"].alignment = wrap
    ws2.column_dimensions["A"].width = 100

    ws2["A4"] = "オープンイシュー / 確認事項"
    ws2["A4"].font = Font(bold=True)
    if spec.open_issues:
        for i, issue in enumerate(spec.open_issues, start=5):
            ws2.cell(row=i, column=1, value=f"- {issue}").alignment = wrap
    else:
        ws2["A5"] = "- なし"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
